#!/usr/bin/env python3
"""PP-X briefing — data charts (PNG) + native PPT text/tables only.

No boxed diagram images. Journal-style plots from numbers only.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

OUT = Path("/Users/baghyeongbae/Desktop/연구/pp_pae_total/output/PP_Research_Detailed_v2.pptx")
ASSETS = Path("/Users/baghyeongbae/Desktop/연구/ppt/pp/_build")
FIGS = ASSETS / "figs"

W, H = 12192000, 6858000
PX = 9525


def px(n: float) -> int:
    return int(n * PX)


C = {
    "ink": RGBColor(0x22, 0x22, 0x22),
    "muted": RGBColor(0x66, 0x66, 0x66),
    "rule": RGBColor(0xE8, 0xE8, 0xE8),
    "blue": RGBColor(0x00, 0x72, 0xB2),
    "orange": RGBColor(0xE6, 0x9F, 0x00),
    "soft": RGBColor(0xF5, 0xF5, 0xF5),
    "soft_blue": RGBColor(0xE8, 0xF1, 0xF8),
    "soft_orange": RGBColor(0xFD, 0xF4, 0xE3),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
    "red": RGBColor(0x9B, 0x1C, 0x1C),
}


def set_run(run, size=18, color=None, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Arial"
    run.font.color.rgb = color or C["ink"]
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        from lxml import etree

        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", "AppleGothic")


def add_text(slide, left, top, width, height, text, size=18, color=None, bold=False, align="left"):
    box = slide.shapes.add_textbox(px(left), px(top), px(width), px(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        run = p.add_run()
        run.text = line
        set_run(run, size=size, color=color or C["ink"], bold=bold)
    return box


def rect(slide, left, top, width, height, fill=None, line=None, rounded=False):
    st = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    sh = slide.shapes.add_shape(st, px(left), px(top), px(width), px(height))
    sh.line.fill.background()
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is not None:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    if rounded:
        try:
            sh.adjustments[0] = 0.12
        except Exception:
            pass
    return sh


def down_arrow(slide, left, top, width, height, fill):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.DOWN_ARROW, px(left), px(top), px(width), px(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def right_arrow(slide, left, top, width, height, fill):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, px(left), px(top), px(width), px(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def vline(slide, x, y1, y2, color):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x), px(y1), px(x), px(y2))
    line.line.color.rgb = color
    line.line.width = Pt(1.25)
    return line


def hline(slide, x1, x2, y, color):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y), px(x2), px(y))
    line.line.color.rgb = color
    line.line.width = Pt(1.25)
    return line


def fill_shape_text(sh, text, size=14, color=None, bold=False, align="center"):
    tf = sh.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.clear()
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    run = p.add_run()
    run.text = text
    set_run(run, size=size, color=color or C["ink"], bold=bold)
    tf.auto_size = None
    sh.word_wrap = True


def pic(slide, name, left, top, width, height):
    path = FIGS / name
    if not path.exists():
        add_text(slide, left, top + 20, width, 24, f"missing: {name}", 12, C["muted"])
        return False
    slide.shapes.add_picture(str(path), px(left), px(top), px(width), px(height))
    return True


def add_table(slide, left, top, width, height, headers, rows, *, font_size=11, highlight_last=False, red_cols=None):
    red_cols = red_cols or set()
    table = slide.shapes.add_table(1 + len(rows), len(headers), px(left), px(top), px(width), px(height)).table
    for j in range(len(headers)):
        table.columns[j].width = px(width / len(headers))

    def cell(r, c, text, *, fill=None, color=None, bold=False, align="center"):
        cl = table.cell(r, c)
        cl.text = ""
        p = cl.text_frame.paragraphs[0]
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[align]
        run = p.add_run()
        run.text = str(text)
        set_run(run, size=font_size, color=color or C["ink"], bold=bold)
        cl.vertical_anchor = MSO_ANCHOR.MIDDLE
        if fill is not None:
            cl.fill.solid()
            cl.fill.fore_color.rgb = fill
        else:
            cl.fill.background()

    for j, h in enumerate(headers):
        cell(0, j, h, fill=C["ink"], color=C["white"], bold=True)
    for i, row in enumerate(rows):
        last = highlight_last and i == len(rows) - 1
        bg = C["soft_blue"] if last else (C["soft"] if i % 2 else C["white"])
        for j, val in enumerate(row):
            col = C["blue"] if last else (C["red"] if j in red_cols else C["ink"])
            cell(i + 1, j, val, fill=bg, color=col, bold=last or j in red_cols or j == 0, align="left" if j == 0 else "center")
    return table


def head(slide, title, sub=""):
    rect(slide, 0, 0, 1280, 3, C["blue"])
    add_text(slide, 48, 22, 1180, 32, title, 22, C["ink"], True)
    if sub:
        add_text(slide, 48, 54, 1180, 20, sub, 12, C["muted"])


def foot(slide, n, total):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(48), px(680), px(1232), px(680))
    line.line.color.rgb = C["rule"]
    line.line.width = Pt(0.75)
    add_text(slide, 48, 686, 800, 14, "SPS Lab  ·  PP-X paper main  ·  frozen method v1", 9, C["muted"])
    add_text(slide, 1100, 684, 100, 16, f"{n}/{total}", 10, C["muted"], True, "right")


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C["white"]
    bg.line.fill.background()
    return s


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    n = 0
    TOTAL = 35

    def p():
        nonlocal n
        n += 1
        return n

    # 1 Cover
    s = blank(prs)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    bg = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C["white"]
    bg.line.fill.background()
    frame = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, px(40), px(32), px(1200), px(656))
    frame.fill.background()
    frame.line.color.rgb = RGBColor(0x22, 0x22, 0x22)
    frame.line.width = Pt(1.25)
    logo = ASSETS / "sps_lab_logo.png"
    if logo.exists():
        s.shapes.add_picture(str(logo), px(860), px(58), px(320), px(64))
    add_text(s, 80, 220, 1120, 40, "PP-X: Prior-Transferability Adaptive Extrapolation", 28, C["ink"], True, "center")
    add_text(s, 80, 268, 1120, 36, "for Robust Prediction Beyond Observed Support", 20, C["ink"], True, "center")
    add_text(s, 80, 330, 1120, 28, "관측된 support 밖에서의 강건 예측", 15, C["muted"], False, "center")
    add_text(s, 80, 390, 1120, 28, "validation-approved prior-residual framework  ·  frozen method v1", 16, C["blue"], True, "center")
    add_text(s, 80, 470, 1120, 24, "Smart Production Systems Lab.  ·  박사과정 박진서", 14, C["ink"], False, "center")
    add_text(s, 80, 520, 1120, 22, "2026.09.11", 13, C["muted"], False, "center")
    p()

    # 2 Problem + background
    s = blank(prs)
    head(s, "문제 · 연구 배경", "밖에서는 데이터가 답을 정하지 못한다.  정당화되는 prior에 맞춰 고른다.")
    pic(s, "nn_vs_ppx_curves.png", 36, 86, 640, 400)
    add_text(s, 36, 492, 640, 36, "점선 왼쪽은 학습, 오른쪽은 그 밖. 제약 없는 NN은 열화 중에도 예측이 되살아난다.", 12, C["muted"])

    rect(s, 696, 86, 548, 132, C["soft"], C["rule"], True)
    add_text(s, 712, 94, 516, 22, "왜 외삽인가", 13, C["ink"], True)
    add_text(s, 712, 118, 516, 90, "학습 분포 안에서는 잘 맞혀도, 관측되지 않은 영역에서는 성능이 급히 떨어질 수 있다. 열화·RUL·피로균열은 실제 시점이 학습 support 밖인 경우가 많아, 보간보다 밖에서의 예측이 중요하다.", 12, C["ink"])

    rect(s, 696, 230, 548, 132, C["soft_orange"], C["orange"], True)
    add_text(s, 712, 238, 516, 22, "prior는 일정하지 않다", 13, C["orange"], True)
    add_text(s, 712, 262, 516, 90, "같은 데이터라도 prior가 바뀌면 밖으로 나가는 곡선이 달라진다. 어떤 문제는 방향·경계·추세만 알고, 어떤 문제는 식과 적용 조건까지 있다. 맞지 않는 prior를 강제하면 외삽이 나빠진다.", 12, C["ink"])

    rect(s, 696, 374, 548, 154, C["soft_blue"], C["blue"], True)
    add_text(s, 712, 382, 516, 22, "출발점", 13, C["blue"], True)
    add_text(s, 712, 408, 516, 108, "맞는 답이 하나가 아니다. 밖을 지탱하는 것은 데이터가 아니라 지금 정당화되는 prior다. prior가 많을수록 좋은 것이 아니라, 그 수준과 신뢰성에 맞춰 외삽 전략을 고른다. 다음 장에서 경로를 나눈다.", 12, C["ink"])
    foot(s, p(), TOTAL)

    # Modeling intent / research gap
    s = blank(prs)
    head(s, "왜 PP-X가 필요한가", "외삽에서는 예측기보다 먼저 구조 가정의 사용 자격을 검증해야 한다")
    add_text(s, 48, 88, 1184, 54, "관측 범위를 벗어나면 데이터만으로 예측 곡선의 형태를 정하기 어렵다. 따라서 어떤 구조 가정을 사용할지 먼저 판단해야 한다.", 18, C["blue"], True, "center")
    rect(s, 48, 170, 510, 294, C["soft_orange"], C["orange"], True)
    add_text(s, 72, 188, 462, 28, "기존 외삽의 취약점", 17, C["orange"], True)
    add_text(s, 78, 238, 440, 190, "① support 밖의 tail은 데이터만으로 식별 불가\n\n② prior와 NN의 수정 권한이 불명확\n\n③ test 결과를 본 뒤 구조를 정당화할 위험\n\n④ 가정이 틀려도 예측값을 강제로 출력", 14, C["ink"])
    right_arrow(s, 584, 294, 74, 42, C["muted"])
    rect(s, 682, 170, 550, 294, C["soft_blue"], C["blue"], True)
    add_text(s, 706, 188, 502, 28, "PP-X의 대응", 17, C["blue"], True)
    add_text(s, 712, 238, 476, 190, "① typed admissibility contract 선언\n\n② frozen prior 중심의 residual 권한 제한\n\n③ physical-unit validation evidence로 승인\n\n④ 미승인 시 exact fallback / abstention", 14, C["ink"])
    add_text(s, 48, 500, 1184, 68, "PP-X는 시험 결과를 보고 유리한 모델을 고르지 않는다. 학습과 검증 단계에서 사용할 가정과 실행 경로를 정하고, 시험 단계에서는 이를 그대로 적용한다.", 16, C["ink"], True, "center")
    add_text(s, 48, 586, 1184, 30, "가정 선언  ·  제한된 residual 학습  ·  검증자료 기반 승인  ·  미승인 시 대체 경로 적용", 13, C["muted"], False, "center")
    foot(s, p(), TOTAL)

    # Contributions and evidence mapping
    s = blank(prs)
    head(s, "본 연구의 기여", "PP-X는 구조 가정의 선언, 학습 범위, 승인 기준을 하나의 절차로 정리한다")
    cards = [
        (
            "01  DECLARE",
            "Typed contract",
            "허용할 prior · tail · 단위 · fallback을 배포 전에 선언",
            "common backbone 6/12\nunit-risk gate 8/12",
            "반례  RWTH dual-scale · MICH fixed-bound",
        ),
        (
            "02  CONTROL",
            "Prior-preserving residual",
            "frozen prior 주변에서 source가 지지하는 만큼만 NN이 수정",
            "direct 17/25, p=.0028\nbound violation 0 / 17,645",
            "미확증  trainable p=.071 · unbounded p=.578",
        ),
        (
            "03  DECIDE",
            "Approval / fallback",
            "물리 단위 증거가 있을 때만 executor를 승인하고 동결",
            "Sun +.221 · MICH +.283\nDS03 route PASS",
            "한계  DS03 .8818 < Engression .9013",
        ),
    ]
    for i, (step, title, body, evidence, limit) in enumerate(cards):
        x = 32 + i * 408
        rect(s, x, 106, 384, 420, C["white"], C["blue"] if i < 2 else C["orange"], True)
        add_text(s, x + 22, 126, 340, 22, step, 12, C["blue"] if i < 2 else C["orange"], True)
        add_text(s, x + 22, 164, 340, 30, title, 18, C["ink"], True)
        add_text(s, x + 22, 216, 340, 76, body, 13, C["ink"])
        rect(s, x + 22, 312, 340, 92, C["soft_blue"])
        add_text(s, x + 36, 324, 312, 66, evidence, 12, C["blue"], True, "center")
        add_text(s, x + 22, 432, 340, 64, limit, 11, C["muted"], False, "center")
    add_text(s, 48, 566, 1184, 46, "개별 모듈보다 중요한 기여는 구조 가정의 선언부터 승인 또는 거절까지를 재현 가능한 절차로 정식화했다는 점이다.", 15, C["blue"], True, "center")
    foot(s, p(), TOTAL)

    # Novelty positioning
    s = blank(prs)
    head(s, "기존 방법과의 차이", "구조 가정의 허용 범위와 승인 절차를 명시했다는 점에서 기존 방법과 구분된다")
    rect(s, 48, 94, 1184, 70, C["soft_blue"], C["blue"], True)
    add_text(s, 72, 111, 1136, 36, "Typed contract · frozen prior 주변의 residual 학습 · validation-only 승인 · 사전 지정 fallback", 16, C["blue"], True, "center")
    comparisons = [
        ("PINN", "완전한 식을 요구하지 않고 typed partial prior의 사용 자격을 검증"),
        ("Prior-residual", "correction을 더하는 데서 끝나지 않고 수정 권한과 승인 규칙을 명시"),
        ("MoE", "test sample별 routing 없이 배포 전에 executor 또는 fallback을 하나로 동결"),
        ("Engression", "generic predictive distribution보다 승인된 structural tail 보존에 초점"),
        ("V-REx", "domain-risk invariance가 아니라 state-level strict-tail과 prior authority를 통제"),
    ]
    for i, (name, diff) in enumerate(comparisons):
        y = 190 + i * 70
        rect(s, 48, y, 174, 48, C["ink"], None, True)
        add_text(s, 58, y + 13, 154, 22, name, 12, C["white"], True, "center")
        rect(s, 236, y, 996, 48, C["soft"] if i % 2 else C["white"], C["rule"], True)
        add_text(s, 256, y + 12, 956, 24, diff, 12, C["ink"])
    add_text(s, 48, 560, 1184, 52, "개별 모듈의 최초성이나 모든 외삽 문제에서의 우월성은 주장하지 않는다.", 13, C["red"], True, "center")
    foot(s, p(), TOTAL)

    # Claim / evidence / reviewer defense
    s = blank(prs)
    head(s, "결과의 해석 범위", "회고 분석, 동일예산 재검증, prospective 평가를 구분해 해석한다")
    tiers = [
        ("TIER A", "회고적 일관성", "9/9  ·  p=.00390625\n77 units  ·  RMSE −33.8%\nCI 15.8–48.6%", "비교모델 예산이 일부 이질적"),
        ("TIER B", "동일예산 재검증", "9×8×30 candidates×5 refit\n8/9  ·  p=.0391\nVirkler −.002", "5 seeds는 독립 표본이 아님"),
        ("TIER C", "Prospective truth test", "route-selection PASS\nfallback .8818\nEngression .9013", "predictive superiority FAIL"),
    ]
    for i, (tier, claim, evidence, limit) in enumerate(tiers):
        x = 30 + i * 410
        edge = C["blue"] if i < 2 else C["orange"]
        rect(s, x, 106, 390, 334, C["white"], edge, True)
        add_text(s, x + 22, 126, 346, 22, tier, 12, edge, True)
        add_text(s, x + 22, 168, 346, 30, claim, 18, C["ink"], True)
        rect(s, x + 22, 222, 346, 118, C["soft_blue"] if i < 2 else C["soft_orange"])
        add_text(s, x + 38, 240, 314, 84, evidence, 13, C["ink"], True, "center")
        add_text(s, x + 22, 370, 346, 42, limit, 11, C["muted"], False, "center")
    rect(s, 48, 474, 1184, 90, C["soft_orange"], C["orange"], True)
    add_text(s, 66, 491, 1148, 56, "PP-X의 구조적 필요성과 회고 분석에서의 일관성은 확인되었다. 다만 독립 prospective 평가에서의 예측 우월성은 아직 확증되지 않았다.", 15, C["ink"], True, "center")
    add_text(s, 48, 594, 1184, 24, "논문의 주장 범위: strict-tail 외삽에서의 구조 검증과 조건부 실행. 보편적 SOTA나 안전 보장은 포함하지 않는다.", 12, C["red"], True, "center")
    foot(s, p(), TOTAL)

    # 4 Research route — full diagram
    s = blank(prs)
    head(s, "연구 루트", "그래서 후보식이 정당화되는지에 따라 경로를 나눈다")

    chips = [("관측", 200), ("경계", 510), ("도메인 지식", 820)]
    for label, x in chips:
        ch = rect(s, x, 92, 260, 40, C["soft"], C["ink"], True)
        fill_shape_text(ch, label, 14, C["ink"], True)
        vline(s, x + 130, 132, 156, C["rule"])
    hline(s, 330, 950, 156, C["rule"])
    vline(s, 640, 156, 170, C["rule"])
    down_arrow(s, 624, 168, 32, 12, C["muted"])

    dec = rect(s, 310, 184, 660, 48, C["ink"], None, True)
    fill_shape_text(dec, "후보식이 정당화되는가?", 17, C["white"], True)

    vline(s, 340, 232, 248, C["rule"])
    vline(s, 940, 232, 248, C["rule"])
    hline(s, 340, 940, 248, C["rule"])
    down_arrow(s, 324, 248, 32, 16, C["blue"])
    down_arrow(s, 924, 248, 32, 16, C["orange"])
    add_text(s, 70, 228, 120, 20, "NO", 12, C["blue"], True, "right")
    add_text(s, 1090, 228, 120, 20, "YES", 12, C["orange"], True, "left")

    rect(s, 70, 272, 540, 200, C["soft_blue"], C["blue"], True)
    rect(s, 70, 272, 8, 200, C["blue"])
    badge = rect(s, 456, 286, 130, 24, C["blue"], None, True)
    fill_shape_text(badge, "이번 발표", 10, C["white"], True)
    add_text(s, 98, 284, 340, 32, "PP-X", 24, C["blue"], True)
    add_text(s, 98, 322, 480, 20, "weak / partial structural prior", 12, C["muted"])
    add_text(s, 98, 352, 490, 22, "typed prior-residual core", 15, C["ink"])
    add_text(s, 98, 384, 490, 22, "+ evidence-selected executor", 15, C["ink"])
    add_text(s, 98, 426, 490, 22, "validation-approved prior-residual  ·  frozen method v1", 14, C["blue"], True)

    rect(s, 670, 272, 540, 200, C["soft_orange"], C["orange"], True)
    rect(s, 670, 272, 8, 200, C["orange"])
    nxt = rect(s, 1056, 286, 130, 24, C["orange"], None, True)
    fill_shape_text(nxt, "future program", 10, C["white"], True)
    add_text(s, 698, 284, 340, 32, "PAE", 24, C["orange"], True)
    add_text(s, 698, 322, 480, 20, "explicit equation-card compiler", 12, C["muted"])
    add_text(s, 698, 352, 490, 22, "허용된 식 후보 생성 + 제한 NN", 15, C["ink"])
    add_text(s, 698, 384, 490, 22, "이득 없으면 PP-X로 되돌림", 15, C["ink"])
    add_text(s, 698, 426, 490, 22, "LLM · 온톨로지 · source gate", 14, C["orange"], True)

    vline(s, 340, 472, 498, C["rule"])
    vline(s, 940, 472, 498, C["rule"])
    hline(s, 340, 940, 498, C["rule"])
    down_arrow(s, 624, 498, 32, 12, C["muted"])

    end = rect(s, 220, 518, 840, 44, C["soft"], C["ink"], True)
    fill_shape_text(end, "Assurance    ·    믿기  /  보류  /  거절    →    박사논문에서 두 경로 통합", 14, C["ink"], True)
    add_text(s, 70, 578, 1140, 24, "오늘은 왼쪽만 간다.  PAE와 Assurance는 지도에만 찍는다.", 13, C["muted"], False, "center")
    foot(s, p(), TOTAL)

    # PP-X at a glance — concise definition before contributions and details.
    s = blank(prs)
    head(s, "PP-X 개요", "검증자료가 지지하는 구조만 사용하고, 그렇지 않으면 미리 정한 대체 경로를 적용한다")
    pic(s, "ppx_framework.png", 58, 92, 1164, 382)
    rect(s, 58, 500, 1164, 82, C["soft_blue"], C["blue"], True)
    add_text(
        s,
        82,
        513,
        1116,
        54,
        "PP-X는 시험 표본마다 모델을 바꾸지 않는다. 학습 전에 구조 가정과 대체 경로를 선언하고, 검증자료로 사용할 실행 구조를 결정한 뒤 시험 단계에서는 이를 고정한다.",
        14,
        C["ink"],
        True,
        "center",
    )
    add_text(s, 58, 606, 1164, 24, "승인된 경우 해당 executor를 사용하고, 승인되지 않으면 사전 지정 fallback을 사용하거나 예측을 보류한다.", 13, C["blue"], True, "center")
    foot(s, p(), TOTAL)

    # 4 Method — final PP-X boundary
    s = blank(prs)
    head(s, "방법 — Algorithm 1: Declare → Learn → Approve → Decline", "test-independent structural assumption approval · frozen method v1")
    stages = [
        ("1 DECLARE", "typed contract\nstate · tail · units · prior · fallback"),
        ("2 LEARN", "frozen prior 주변\nsource-supported bounded residual"),
        ("3 APPROVE", "validation-only\nphysical-unit gain · unit-risk"),
        ("4 DECLINE", "승인 실패\nexact fallback / abstention"),
    ]
    for i, (label, body) in enumerate(stages):
        x = 28 + i * 310
        sh = rect(s, x, 84, 278, 116, C["soft_blue"] if i < 3 else C["soft_orange"], C["blue"] if i < 3 else C["orange"], True)
        add_text(s, x + 12, 94, 254, 20, label, 13, C["blue"] if i < 3 else C["orange"], True, "center")
        add_text(s, x + 12, 122, 254, 62, body, 11, C["ink"], False, "center")
        if i < 3:
            right_arrow(s, x + 282, 130, 24, 20, C["muted"])
    pic(s, "ppx_core.png", 20, 218, 760, 300)
    add_table(
        s,
        790,
        220,
        448,
        132,
        ["단계", "frozen method v1"],
        [
            ["Authority", "prior 중심 residual 수정 권한 제한"],
            ["Selection", "test-independent · sample-wise routing 없음"],
        ],
        font_size=11,
    )
    add_table(
        s,
        790,
        366,
        448,
        154,
        ["Safety gate (val만)", "승인 조건"],
        [
            ["상대 이득", "same-split fallback보다 RMSE 개선"],
            ["유닛 증거", "unit-bootstrap 95% CI 하한 > 0"],
            ["승인", "validation evidence가 지지한 executor"],
            ["거절", "frozen fallback 또는 abstention"],
        ],
        font_size=11,
    )
    add_text(
        s,
        20,
        530,
        760,
        52,
        "PP-X top-level은 모든 데이터에서 하나의 동일 network를 강제한다는 뜻이 아니다. core와 executor 후보를 validation evidence로 선택한다.\n"
        "trust=0 exact matched MLP는 역사적 nested stress-test의 한 fallback 구현이다. 최종 방법은 선언된 frozen fallback을 쓰거나 예측을 abstain한다.",
        12,
        C["ink"],
    )
    add_text(s, 20, 600, 1210, 22, "Algorithm 1은 test 예측을 보고 혼합하지 않는다. 승인된 한 경로만 frozen forward; 실패는 exact fallback 또는 abstention.", 12, C["muted"])
    foot(s, p(), TOTAL)

    # 5 How the evaluation interval is defined
    s = blank(prs)
    head(s, "외삽 구간을 어떻게 정하는가", "학습이 본 건강 범위보다 더 진행된 관측만 점수로 친다")
    pic(s, "extrapolation_cut.png", 36, 86, 720, 350)
    add_table(
        s,
        770,
        88,
        462,
        350,
        ["절차", "내용"],
        [
            ["1", "경계, 유닛, 시간, 입력, 외삽 좌표를 먼저 고정"],
            ["2", "학습·검증·시험에 같은 유닛을 넣지 않음"],
            ["3", "학습에는 본 적 있는 건강 구간만 사용"],
            ["4", "검증·시험에는 그 범위 밖 관측만 남김"],
            ["5", "정규화 통계는 학습 분할에서만 계산"],
            ["6", "모델과 executor는 검증에서만 선택"],
            ["7", "고정한 시험 분할은 한 번만 예측"],
        ],
        font_size=11,
    )
    add_text(
        s,
        36,
        452,
        1200,
        100,
        "세로선은 학습에서 관측한 최저 건강이다. 그보다 더 나빠진 구간(오른쪽)만 검증·시험 점수에 넣는다.\n같은 유닛의 미래 값, 최종 수명, 시험 궤적에서 만든 통계는 입력에 쓰지 않는다. 이 조건을 어기면 내삽 평가가 된다.",
        13,
        C["ink"],
    )
    add_text(s, 36, 562, 1200, 28, "자르는 규칙은 같다. 외삽 geometry는 셋마다 다르므로 다음 장에서 1D와 다차원을 구분해 적는다.", 12, C["muted"])
    foot(s, p(), TOTAL)

    # Quantification of the cut
    s = blank(prs)
    head(s, "외삽을 숫자로 확인하는 방법", "1D 열화좌표 기준이다. 전체 feature-space hull 밖이라고 말하지 않는다.")
    pic(s, "hull_distance.png", 20, 82, 520, 400)
    add_table(
        s,
        550,
        82,
        690,
        380,
        ["셋", "1D Hull", "거리", "성격"],
        [
            ["선우다", "100%", "4.42", "미지 셀 · 늦은 건강. 매우 명확"],
            ["아헨", "100%", "2.87", "미지 셀 · 늦은 건강. 매우 명확"],
            ["MIT 2019", "100%", "5.55", "1D 건강 support 밖"],
            ["미시간", "100%", "3.43", "미지 셀 · 늦은 건강"],
            ["NASA 실험셀", "100%", "2.01", "미지 셀 · 건강 tail"],
            ["화중 배터리", "100%", "1.61", "미지 유닛 · 충전법/말기. 1D 기준"],
            ["알루미늄 균열", "100%", "1.62", "미지 시편 · 균열 tail"],
            ["MIT 배치2", "100%", "1.90", "1D 100% · PCA2 0% · PCA3 99.2%"],
            ["항공기 엔진", "100%", "0.23", "1D 100% · PCA2 0% · PCA3 83.6% · 조건"],
        ],
        font_size=10,
    )
    add_text(
        s,
        20,
        498,
        1240,
        110,
        "오늘 그림에는 1D ordered-coordinate 기준의 엄격한 외삽만 남긴다. XJTU·FEMTO·NASA milling은 도메인/재료 transfer라 이 표에 없다.\n1D Hull-out 100%는 ‘선언한 건강·균열·조건 축’ 밖이지, 전체 특징공간 convex hull 밖이 아니다. 배치2와 엔진은 PCA 2D에서 0%다.\n엔진은 거리 0.23 SD·Target-out 0%라 먼 수명 외삽이 아니라 운전조건(regime) 외삽이다.",
        12,
        C["ink"],
    )
    foot(s, p(), TOTAL)

    # Protocol — after method, before scores
    s = blank(prs)
    head(s, "프로토콜", "점수는 이 범위 안에서만 읽는다  ·  적합도는 허가증이지 성적표가 아니다")
    add_table(
        s,
        48,
        100,
        1184,
        250,
        ["규칙", "의미", "깨지면"],
        [
            ["unit-disjoint", "train / val / test 셀이 겹치지 않음", "같은 유닛 누수"],
            ["hull-out", "선언한 1D 열화좌표에서 test가 train 밖", "내삽으로 바뀜"],
            ["val-only", "게이트·보정을 val에서만 고름", "test로 튜닝"],
            ["Pass / Weak / Fail", "운행 허가. Fail이면 예측을 옮기지 않음", "억지 점수"],
        ],
        font_size=13,
    )
    add_table(
        s,
        48,
        380,
        580,
        220,
        ["In scope", "예"],
        [
            ["1D late-tail", "Sun · RWTH · MICH · MATR2019 · NASA · Virkler · HUST"],
            ["1D + 다차원 설명 필요", "MATRb2 · N-CMAPSS (PCA2는 0%)"],
            ["executor 선택", "train/val · group-LOO only"],
        ],
        font_size=12,
    )
    add_table(
        s,
        652,
        380,
        580,
        220,
        ["오늘 표에 없음", "이유"],
        [
            ["전체 feature hull", "1D 100% ≠ 다차원 hull 밖"],
            ["도메인 / 재료 transfer", "조건·끝점·재질 이동은 다른 종류"],
            ["C-MAPSS FD002/004", "조건 hull은 강하나 이번 주표 밖"],
        ],
        font_size=12,
    )
    foot(s, p(), TOTAL)

    # 6 Main results — retrospective PP-X portfolio
    s = blank(prs)
    head(s, "주 결과 — PP-X 9-setting portfolio", "retrospective concept-aligned evidence · 확증 cohort 아님 · PAE 없음")
    pic(s, "ppx_portfolio.png", 30, 82, 680, 390)
    add_table(
        s,
        720,
        88,
        512,
        384,
        ["셋", "Executor", "R²"],
        [
            ["화중 배터리", "transport", "0.958"],
            ["선우다 상용셀", "fixed BQ", "0.939"],
            ["항공기 엔진", "multiscale", "0.937"],
            ["알루미늄 균열", "gated residual", "0.888"],
            ["아헨 배터리", "fixed BQ", "0.878"],
            ["MIT 배치2", "decay+transport", "0.862"],
            ["미시간 배터리", "dual-scale", "0.751"],
            ["NASA 실험셀", "multiscale", "0.584"],
            ["MIT 2019", "cal. latent", "0.466"],
        ],
        font_size=10,
    )
    rect(s, 36, 488, 400, 168, C["soft_blue"], C["blue"], True)
    rect(s, 36, 488, 8, 168, C["blue"])
    add_text(s, 56, 498, 360, 22, "Li 배터리", 13, C["blue"], True)
    add_text(
        s,
        56,
        524,
        360,
        120,
        "화중·선우다·아헨·미시간 리튬셀\nMIT 수명 벤치 (2019 / 배치2)\nNASA 실험용 18650  (로켓 아님)",
        11,
        C["ink"],
    )
    rect(s, 452, 488, 380, 168, C["soft"], C["ink"], True)
    rect(s, 452, 488, 8, 168, C["ink"])
    add_text(s, 472, 498, 340, 22, "균열", 13, C["ink"], True)
    add_text(
        s,
        472,
        524,
        340,
        110,
        "알루미늄 판에 금이 감\n항공기 재료 피로실험\n배터리 아님",
        12,
        C["ink"],
    )
    rect(s, 848, 488, 396, 168, C["soft_orange"], C["orange"], True)
    rect(s, 848, 488, 8, 168, C["orange"])
    add_text(s, 868, 498, 356, 22, "엔진", 13, C["orange"], True)
    add_text(
        s,
        868,
        524,
        356,
        110,
        "항공기 엔진 시뮬레이터\n처음 보는 비행조건\n실제 비행기 데이터가 아님",
        12,
        C["ink"],
    )
    foot(s, p(), TOTAL)

    # Equal-budget retrospective evidence
    s = blank(prs)
    head(s, "동일예산 · retrospective 평가", "uniformly tuned: 9 settings × 8 baselines × 30 candidates × 5 refit")
    add_table(
        s,
        48,
        92,
        1184,
        180,
        ["설계", "값", "의미"],
        [
            ["Settings", "9", "동일한 retrospective 평가 단위"],
            ["Baselines", "8", "각 setting에서 같은 예산으로 비교"],
            ["Candidates", "30", "방법별 후보 수 고정"],
            ["Refit", "5", "선택 후 동일 횟수 재적합"],
        ],
        font_size=13,
    )
    rect(s, 48, 304, 360, 220, C["soft_blue"], C["blue"], True)
    add_text(s, 68, 320, 320, 28, "전체 판정", 18, C["blue"], True)
    add_text(s, 68, 366, 320, 112, "PP-X  8 / 9\n양측 sign p=.0391\n\n동일예산에서 우세 방향", 16, C["ink"], False, "center")

    rect(s, 460, 304, 360, 220, C["soft_orange"], C["orange"], True)
    add_text(s, 480, 320, 320, 28, "유일한 미승리", 18, C["orange"], True)
    add_text(s, 480, 366, 320, 112, "Virkler  −0.002\n\n차이는 작지만\n승리로 세지 않는다", 16, C["ink"], False, "center")

    rect(s, 872, 304, 360, 220, C["soft"], C["ink"], True)
    add_text(s, 892, 320, 320, 28, "해석 경계", 18, C["ink"], True)
    add_text(s, 892, 360, 320, 126, "후보 30회 = 방법별 탐색예산\n\nPP-X 구조개발 전체가\n30회였다는 뜻 아님", 14, C["ink"], False, "center")

    add_text(s, 48, 552, 1184, 56, "해석  동일한 후보·refit 예산에서 우세 방향을 재검사한 retrospective benchmark다. 5 seeds는 독립 표본 아님. sign test의 n은 9 settings다.", 13, C["ink"])
    add_text(s, 48, 616, 1184, 22, "Tier A mixed strongest portfolio 및 Tier C DS03 prospective와 결합하거나 pooled하지 않는다.", 12, C["muted"], True)
    foot(s, p(), TOTAL)

    # DS03 prospective truth test
    s = blank(prs)
    head(s, "DS03 prospective truth test", "Tier C · route-selection PASS와 predictive superiority FAIL을 분리한다")
    add_table(
        s, 48, 96, 1184, 246,
        ["모델 / 경로", "R²", "사전 판정", "해석"],
        [
            ["PP-X selected fallback", ".8818", "선택", "route-selection PASS"],
            ["PP-X basic", ".832", "미선택", "fallback보다 낮음"],
            ["PP-X multiscale", ".869", "미선택", "fallback보다 낮음"],
            ["Engression", ".9013", "비교", "predictive superiority FAIL"],
            ["CRT / GCIE", ".8851 / .8850 (nested .8727)", "rejected", "Algorithm 1 미포함"],
        ], font_size=12,
    )
    rect(s, 48, 380, 570, 142, C["soft_blue"], C["blue"], True)
    add_text(s, 68, 394, 530, 24, "성공한 주장", 16, C["blue"], True)
    add_text(s, 68, 430, 530, 70, "test를 보기 전에 fallback 경로를 선택했고,\n대안 PP-X 경로보다 높았다.", 14, C["ink"], False, "center")
    rect(s, 662, 380, 570, 142, C["soft_orange"], C["orange"], True)
    add_text(s, 682, 394, 530, 24, "실패한 주장", 16, C["orange"], True)
    add_text(s, 682, 430, 530, 70, "Engression .9013을 넘지 못했다.\nprospective predictive superiority를 주장하지 않는다.", 14, C["ink"], False, "center")
    add_text(s, 48, 562, 1184, 50, "N-CMAPSS .937은 retrospective portfolio split이며 DS03가 아니다. Tier A/B 수치와 DS03를 합산하지 않는다.", 13, C["red"], True, "center")
    foot(s, p(), TOTAL)

    # Final model boundary
    s = blank(prs)
    head(s, "최종 모델 경계", "PP-X가 top-level · executor와 rejected 개발안을 구분한다")
    add_table(
        s,
        48,
        92,
        1184,
        330,
        ["구성", "역할 / 결과", "최종 판정"],
        [
            ["PP-X", "validation-approved prior-residual framework", "paper main · frozen method v1"],
            ["SAAR", "support-aware affine-residual mechanism", "historical alias / core architecture"],
            ["CCMR v2.2", "trajectory-domain 실행", "executor only · 메인 아님"],
            ["CRT", "DS03 .8851 < Engression .9013", "rejected · Algorithm 1 미포함"],
            ["GCIE", ".8850 · nested .8727", "rejected · Algorithm 1 미포함"],
        ],
        font_size=12,
    )
    rect(s, 48, 454, 570, 126, C["soft_blue"], C["blue"], True)
    add_text(s, 68, 466, 530, 24, "승인될 때", 15, C["blue"], True)
    add_text(s, 68, 500, 530, 64, "validation evidence가 지지한 core / domain executor만 실행", 14, C["ink"], False, "center")
    rect(s, 662, 454, 570, 126, C["soft_orange"], C["orange"], True)
    add_text(s, 682, 466, 530, 24, "prior가 거절될 때", 15, C["orange"], True)
    add_text(s, 682, 500, 530, 64, "사전 고정 fallback으로 전환하거나 abstention · test로 재선택 금지", 14, C["ink"], False, "center")
    add_text(s, 48, 600, 1184, 30, "PAE는 future program일 뿐 paper main이 아니다. CRT/GCIE는 최종 Algorithm 1에 넣지 않는다.", 13, C["muted"], True, "center")
    foot(s, p(), TOTAL)

    # Historical v1.1 model explanation — retain legacy ablation evidence.
    s = blank(prs)
    head(s, "[Appendix] 역사적 메커니즘 — trust stress-test", "개발 당시 v1.1 nested network · 최종 PP-X 전체 구조가 아님")
    add_text(s, 48, 78, 1184, 30, "이 장의 동일-network 설계는 trust와 exact fallback을 검증한 historical mechanism/stress-test다.", 15, C["ink"], True)
    rect(s, 48, 134, 340, 330, C["soft_orange"], C["orange"], True)
    add_text(s, 70, 150, 296, 26, "① 후보 생성", 18, C["orange"], True)
    add_text(s, 70, 198, 296, 220, "trust > 0\n동결 affine prior\n+ nonlinear residual\n\ntrust = 0\n동일 초기화·optimizer의\nmatched MLP", 15, C["ink"], False, "center")
    add_text(s, 398, 270, 42, 40, "→", 25, C["muted"], True, "center")
    rect(s, 446, 134, 340, 330, C["soft_blue"], C["blue"], True)
    add_text(s, 468, 150, 296, 26, "② validation 승인", 18, C["blue"], True)
    add_text(s, 468, 198, 296, 220, "MLP 대비 RMSE\n상대 2% 이상 개선\n+\nphysical-unit bootstrap\n95% CI 하한 > 0\n+\nvalidation unit ≥ 3", 14, C["ink"], False, "center")
    add_text(s, 796, 270, 42, 40, "→", 25, C["muted"], True, "center")
    rect(s, 844, 134, 388, 330, C["soft"], C["ink"], True)
    add_text(s, 866, 150, 344, 26, "③ 한 경로만 출력", 18, C["ink"], True)
    add_text(s, 866, 198, 344, 220, "통과  prior trust 유지\n\n실패  trust = 0\nexact matched MLP\n\n※ test 예측을 보고\n혼합하지 않는다", 15, C["ink"], False, "center")
    rect(s, 48, 500, 1184, 76, C["ink"], None, True)
    add_text(s, 70, 517, 1140, 46, "historical v1.0 = core 구조 점검   ·   historical v1.1 = trust/fallback stress-test", 15, C["white"], True, "center")
    add_text(s, 48, 590, 1184, 26, "주의  Stanford 결과를 본 뒤 만든 post-test development다. retrospective portfolio와 final Algorithm 1을 소급 대체하지 않는다.", 12, C["red"], True)
    foot(s, p(), TOTAL)

    # Historical v1.1 trust/architecture stress-test
    s = blank(prs)
    head(s, "[Appendix] Historical stress-test — trust · architecture", "개발 당시 v1.1 · 2고호트 × 6 trust × 4 architecture × 5 seeds = 240 fits")
    pic(s, "v11_complete_trust.png", 28, 82, 760, 400)
    add_table(
        s,
        808,
        92,
        420,
        350,
        ["Cohort", "Val-best trust", "Test-best*"],
        [
            ["Stanford", "0", ".20"],
            ["ISU 250mAh", ".02", ".10"],
            ["Grid", "w 16/32", "lr .0005/.001"],
            ["Seeds", "42–46", "5/arm"],
        ],
        font_size=12,
        red_cols={2},
    )
    add_text(s, 808, 458, 420, 42, "* test-best는 설명용 사후 dose-response이며 선택에 사용하지 않음", 11, C["red"], True)
    add_text(
        s,
        48,
        510,
        1184,
        82,
        "Stanford  validation은 trust=0을 선호하지만 test dose 최고는 .20(R² .094).  ISU는 validation .02, test dose 최고 .10(R² .557).\n"
        "결론  trust의 test 최적점은 고호트마다 다르고 validation 최적점과도 다르다. test를 보고 trust를 고르면 누수다.",
        13,
        C["ink"],
    )
    add_text(s, 48, 606, 1184, 22, "historical mechanism evidence only · final PP-X 전체가 하나의 동일 network라는 뜻이 아니다.", 12, C["muted"], True)
    foot(s, p(), TOTAL)

    # Historical v1.1 gate factorial and seed stability
    s = blank(prs)
    head(s, "[Appendix] Historical stress-test — gate 2×2 · seed", "개발 당시 v1.1 · 2% margin on/off × unit-bootstrap on/off · exact fallback")
    pic(s, "v11_complete_gate_seed.png", 28, 82, 760, 390)
    add_table(
        s,
        808,
        92,
        420,
        330,
        ["Cohort", "Val gain", "Boot 95% CI", "Full gate"],
        [
            ["Stanford", "−0.28%", "[−.276,.208]", "reject"],
            ["ISU 250", "+1.65%", "[−.145,.196]", "reject"],
            ["Unit wins", "4/8", "32/45", "—"],
            ["Fallback Δ", "0.0", "0.0", "exact"],
        ],
        font_size=10,
    )
    add_text(s, 808, 444, 420, 54, "각 gate 단독으로도 두 prior를 거절\nFull gate = matched MLP exact", 13, C["blue"], True)
    add_text(
        s,
        48,
        510,
        1184,
        82,
        "비용  ISU always-on test R² .555 → full gate .451. 유효한 약한 prior도 놓친다.\n"
        "Seed  두 고호트 모두 개별 seed 하나는 R²<0. 따라서 ensemble 양수만으로 seed-robust를 주장하지 않는다.",
        13,
        C["ink"],
    )
    add_text(s, 48, 606, 1184, 22, "다음 legacy ablation과 이 두 고호트는 retrospective mechanism evidence다. final Algorithm 1의 동일-network 주장이 아니다.", 12, C["muted"], True)
    foot(s, p(), TOTAL)

    # Ablation — dual-scale is one executor
    s = blank(prs)
    head(s, "[Appendix] Ablation — dual-scale", "executor 하나  ·  모든 데이터에 켜는 기본값이 아니다")
    pic(s, "ablation_panel.png", 24, 82, 1232, 430)
    add_text(
        s,
        48,
        520,
        580,
        90,
        "(a) 결론  고정 bound는 Sun·RWTH에서는 높지만 MICH를 0.468에 묶는다. dual-scale을 켠 뒤에야 MICH가 0.751이 된다.",
        13,
        C["ink"],
    )
    add_text(
        s,
        660,
        520,
        572,
        90,
        "(b) 결론  평균을 조금 낮추고 최저 점수를 올린 선택이다. 데이터별 최고점을 모은 모델이 아니다.",
        13,
        C["ink"],
    )
    foot(s, p(), TOTAL)

    # Component ΔR²
    s = blank(prs)
    head(s, "[Appendix] Ablation — 구성요소", "켠 모델 − matched 제거 arm  ·  ΔR²")
    pic(s, "ablation_delta.png", 36, 88, 760, 400)
    add_table(
        s,
        810,
        88,
        430,
        400,
        ["기능", "대표", "ΔR²"],
        [
            ["Residual", "Sun", "+0.658"],
            ["Residual", "MICH", "+3.811"],
            ["Frozen affine", "MICH", "+0.149"],
            ["Fixed bound", "Sun", "+0.221"],
            ["Fixed bound", "MICH", "−0.291"],
            ["Dual-scale", "MICH", "+0.283"],
            ["Rate hist.", "RWTH", "+1.256"],
            ["Transport", "MATRb2", "+0.187"],
        ],
        font_size=11,
    )
    add_text(s, 36, 520, 760, 70, "결론  residual은 세 배터리에서 모두 이득이다. 고정 bound와 full history는 MICH에서 마이너스다. 그래서 모듈을 쌓지 않고, 근거 있는 executor만 켠다.", 12, C["ink"])
    foot(s, p(), TOTAL)

    # 6-arm + history
    s = blank(prs)
    head(s, "[Appendix] Ablation — 대조군 · history", "(a) matched 6-arm  ·  (b) causal history")
    pic(s, "ablation_arms.png", 20, 88, 630, 400)
    pic(s, "ablation_history.png", 650, 88, 600, 380)
    add_text(s, 48, 500, 1180, 80, "(a) 결론  Affine만, 또는 그냥 NN만으로는 부족하다. 동결 affine 위에 제한 residual을 올린 조합이 세 셋에서 가장 안정하다.\n(b) 결론  속도 이력은 Sun·RWTH에 필요하고, MICH에서는 단순한 margin history가 더 높다(0.715 vs 0.468). 이력을 전역 기본값으로 두지 않는다.", 13, C["ink"])
    foot(s, p(), TOTAL)

    # MATR 2x2 + gate
    s = blank(prs)
    head(s, "[Appendix] Ablation — transport · gate", "optional executor는 항상 켜지 않는다")
    pic(s, "ablation_matr.png", 40, 100, 500, 420)
    pic(s, "stats_gate.png", 560, 100, 680, 400)
    add_text(
        s,
        40,
        518,
        1200,
        80,
        "(a) 결론  MATRb2에서 transport가 주효과(+0.187)다. decay만으로는 +0.002다. 둘을 같이 켜야 0.862다.\n(b) 결론  gate는 성능을 만드는 장치가 아니라 거절 장치다. 13곳 중 3곳만 개선, 10곳은 유지, 악화 0. seed는 5/5일 때만 승인한다.",
        12,
        C["ink"],
    )
    foot(s, p(), TOTAL)

    # Ablation synthesis
    s = blank(prs)
    head(s, "Ablation 종합", "모듈을 쌓지 않는다.  검증에서 이득이 있는 executor만 켠다.")
    add_text(s, 48, 86, 1184, 28, "앞 네 장의 숫자를 한 규칙으로 읽는다. 최종 PP-X는 공통 core + 데이터마다 고른 executor다.", 14, C["ink"])

    rect(s, 48, 124, 380, 360, C["soft_blue"], C["blue"], True)
    rect(s, 48, 124, 8, 360, C["blue"])
    add_text(s, 68, 136, 340, 28, "항상 켠다", 18, C["blue"], True)
    add_text(s, 68, 172, 340, 22, "core", 12, C["muted"])
    add_text(
        s,
        68,
        202,
        340,
        260,
        "동결 affine prior\n+ 제한 residual\n\nAffine만, 또는 NN만으로는\n세 배터리에서 무너진다.\n\nResidual ΔR²\nSun +0.66  ·  MICH +3.81",
        14,
        C["ink"],
    )

    rect(s, 450, 124, 380, 360, C["soft_orange"], C["orange"], True)
    rect(s, 450, 124, 8, 360, C["orange"])
    add_text(s, 470, 136, 340, 28, "근거 있을 때만", 18, C["orange"], True)
    add_text(s, 470, 172, 340, 22, "executor  ·  val-only", 12, C["muted"])
    add_text(
        s,
        470,
        202,
        340,
        260,
        "고정 bound  Sun · RWTH\ndual-scale  MICH만 (+0.28)\n속도 이력  Sun · RWTH\ntransport  HUST · MATRb2\n\ngate  13곳 중 3곳만 승인\n악화 0",
        14,
        C["ink"],
    )

    rect(s, 852, 124, 380, 360, C["soft"], C["ink"], True)
    rect(s, 852, 124, 8, 360, C["ink"])
    add_text(s, 872, 136, 340, 28, "켜면 나빠진다", 18, C["ink"], True)
    add_text(s, 872, 172, 340, 22, "전역 기본값으로 두지 않음", 12, C["muted"])
    add_text(
        s,
        872,
        202,
        340,
        260,
        "고정 bound → MICH −0.29\nfull rate history → MICH −0.25\ndual-scale → Sun −0.005,\nRWTH −0.036\n\n평균을 조금 깎고\n최저점을 살리는 선택은\n데이터별 최고점 모음이 아니다.",
        14,
        C["ink"],
    )

    end = rect(s, 48, 504, 1184, 72, C["ink"], None, True)
    fill_shape_text(end, "읽는 법    core는 고정한다.   executor는 검증 증거가 있을 때만 켠다.   실패하면 safety로 되돌린다.", 15, C["white"], True)
    add_text(s, 48, 586, 1184, 24, "그래서 주표의 executor가 데이터마다 다르다.  한꺼번에 켠 공동 모델이 아니다.", 13, C["muted"])
    foot(s, p(), TOTAL)

    # Competitors — after internal ablation, before formal tests
    s = blank(prs)
    head(s, "비교", "(a) heatmap  ·  (b) PP-X vs TabPFN vs others")
    pic(s, "competitor_bars.png", 16, 72, 1248, 528)
    add_text(s, 40, 608, 1200, 40, "1 PP-X 0.81 · 2 GroupDRO 0.36 · 3 V-REx 0.35.  순위=9곳 평균 R².  강건=9곳에서 양수(9/9, 7/9, 7/9).  MICH TabPFN=동일 202행, v3 CPU, ensemble −1.86.", 13, C["ink"])
    foot(s, p(), TOTAL)

    # Stats then robustness — one evidence block
    s = blank(prs)
    head(s, "통계", "무엇을 검정했는가  ·  파랑 = p<0.05  ·  주황 = 유의 못 함")
    pic(s, "stats_wilcoxon.png", 16, 82, 568, 340)
    add_table(
        s,
        590,
        82,
        650,
        340,
        ["①", "검정", "질문", "결과"],
        [
            ["W", "Wilcoxon", "Sun·RWTH·MICH 25 unit RMSE", "Direct 17/25 p=.003"],
            ["W", "Wilcoxon", "soft / affine 대비", "24/25 · 25/25"],
            ["W", "Wilcoxon", "trainable / unbounded", "p=.071 · .578  못 함"],
            ["U", "Unit wins", "아홉 셋 물리 유닛 77개", "60/77 이김"],
            ["A", "Mixed portfolio sign", "strongest same-split · heterogeneous budget", "9/9 · p=.00390625"],
            ["B", "Equal-budget sign", "uniformly tuned retrospective", "8/9 · p=.0391"],
            ["B", "Bound audit", "residual이 이론 bound를 넘나", "0 / 17,645"],
        ],
        font_size=10,
    )
    rect(s, 28, 432, 300, 200, C["soft_blue"], C["blue"], True)
    add_text(s, 40, 440, 276, 20, "W  Wilcoxon", 12, C["blue"], True)
    add_text(s, 40, 464, 276, 155, "쌍을 이룬 unit RMSE.\nH0: 중앙 차이 = 0.\n파랑만 ‘이겼다’고 말함.\n주황은 이긴 칸이 있어도 유의 아님.", 11, C["ink"])
    rect(s, 340, 432, 300, 200, C["soft"], C["ink"], True)
    add_text(s, 352, 440, 276, 20, "S  Seed binomial", 12, C["ink"], True)
    add_text(s, 352, 464, 276, 155, "5 seed가 affine을 골랐는가.\nH0: 확률 ≤ 0.5.\n5/5만 α=0.05 통과.\n물리 반복이 아님.", 11, C["ink"])
    rect(s, 652, 432, 300, 200, C["soft_orange"], C["orange"], True)
    add_text(s, 664, 440, 276, 20, "D  Domain sign", 12, C["blue"], True)
    add_text(s, 664, 464, 276, 155, "Tier A: 9/9, p=.00390625.\ngeo. RMSE −33.8%\nhierarchical CI 15.8–48.6%.\nTier B: 8/9, p=.0391.\n서로 합치지 않음.", 10, C["ink"])
    rect(s, 964, 432, 276, 200, C["soft"], C["ink"], True)
    add_text(s, 976, 440, 252, 20, "B  Bound audit", 12, C["ink"], True)
    add_text(s, 976, 464, 252, 155, "가설검정이 아니라 제약 감사.\n|ŷ−affine| ≤ margin·B\n17,645점 위반 0.", 11, C["ink"])
    foot(s, p(), TOTAL)

    s = blank(prs)
    head(s, "안정성", "유닛 단위에서 개선이 한쪽으로 몰리지 않았는지 확인한다")
    pic(s, "robustness_panel.png", 40, 80, 1200, 380)
    add_text(
        s,
        48,
        468,
        580,
        120,
        "(a) 결론  아홉 셋 유닛 log-RMSE 비는 평균 0.39, 계층 bootstrap 95% CI [0.16, 0.65]. 선우다·배치2·엔진·아헨은 구간이 0을 넘는다. 화중·균열·NASA·2019는 유닛이 적어 구간이 0을 포함한다.",
        13,
        C["ink"],
    )
    add_text(
        s,
        660,
        468,
        572,
        120,
        "(b) 결론  MICH 시험 8유닛의 개별 R²가 모두 양수다(0.50–0.96). 합친 점수 0.751은 한 유닛에 몰린 값이 아니므로 대표값으로 읽어도 된다. 이 그림은 검정이 아니라 분포 확인이다.",
        13,
        C["ink"],
    )
    add_text(s, 48, 598, 1184, 28, "정리  유닛 안에서는 근거가 있다. 데이터 종류가 세 개뿐이라 분야 전체 유의는 말하지 않는다.", 13, C["muted"])
    foot(s, p(), TOTAL)

    # 10 Failures — tables only
    s = blank(prs)
    head(s, "실패 · 경계", "오늘 표는 1D 엄격한 외삽만.  같은 100%라도 geometry는 다르다.")
    add_table(
        s,
        48,
        110,
        1184,
        230,
        ["셋", "R²", "경계", "읽는 법"],
        [
            ["MICH (base)", "-1.522", "관계 이동 · 보정 꺼짐", "dual-scale 켠 뒤 0.751"],
            ["MATR2019", "0.466", "1D 건강 tail은 맞음", "주표에 남기되 약점"],
            ["N-CMAPSS", "0.937*", "1D 100% · PCA2 0%", "portfolio split. DS03 아님"],
            ["MATRb2", "0.862", "1D 100% · PCA2 0% · PCA3 99.2%", "다차원 geometry를 같이 적음"],
        ],
        font_size=12,
        red_cols={1},
    )
    add_table(
        s,
        48,
        380,
        580,
        220,
        ["Claim", "Detail"],
        [
            ["Model", "frozen method v1 · 표 수치는 retrospective portfolio"],
            ["Scope", "unit-disjoint · hull-out · val-only"],
            ["Not default", "dual-scale · transport · full history"],
            ["Position", "방법론·신뢰성 중심 상위저널"],
        ],
        font_size=12,
    )
    add_table(
        s,
        660,
        380,
        572,
        220,
        ["Not claimed", "Detail"],
        [
            ["Zn / Na", "사후 개발 — untouched 확증 아님"],
            ["도메인 transfer", "오늘 외삽 정의에 넣지 않음"],
            ["PAE 주표", "오늘 점수와 섞지 않음"],
            ["Universal SOTA", "모든 OOD 1등 아님"],
        ],
        font_size=12,
    )
    add_text(s, 48, 616, 1184, 22, "* N-CMAPSS 0.937은 retrospective portfolio split이다. prospective DS03와 동일 split·비교가 아니다.", 11, C["muted"], True)
    foot(s, p(), TOTAL)

    # Locked cohorts already scored — no new download
    s = blank(prs)
    head(s, "[Appendix] 끝난 고호트 — PP-X", "잘된 셋만 PP-X.  MATR 두 곳은 기존 PP-X artifact  ·  Misata는 방금 같은 split으로 재실행")
    add_table(
        s,
        48,
        92,
        1184,
        360,
        ["고호트", "PP-X 경로", "PP-X R²", "판정"],
        [
            ["MATR 2019-01-24", "validation-calibrated latent", "0.466", "양수. 봉인 latent 0.257에서 개선. 주표와 동일"],
            ["MATR batch2", "regime transport", "0.862", "양수. 봉인 0.471/0.523에서 개선. 주표와 동일"],
            ["Misata", "core · val이 unbounded 선택", "0.829", "양수. MLP 0.854. 우월 확증은 여전히 실패"],
        ],
        font_size=13,
        red_cols={2},
    )
    add_text(
        s,
        48,
        470,
        1184,
        90,
        "Misata 새 실행  seed 42–46 ensemble 0.829 (개별 0.816–0.829). executor는 validation에서 bounded가 2%를 못 넘겨 unbounded. 기존 locked MLP와 같은 1298행.\nOxford·MATWI·팬·XJTU는 실패 고호트라 PP-X로 다시 돌리지 않았다.",
        13,
        C["ink"],
    )
    foot(s, p(), TOTAL)

    # Close — forward only, not a repeat of slide 2
    s = blank(prs)
    head(s, "다음", "지도는 루트 장에 있다.  여기서는 앞으로만 말한다.")

    rect(s, 48, 110, 380, 360, C["soft_blue"], C["blue"], True)
    rect(s, 48, 110, 8, 360, C["blue"])
    now = rect(s, 286, 126, 118, 24, C["blue"], None, True)
    fill_shape_text(now, "지금", 11, C["white"], True)
    add_text(s, 72, 126, 200, 32, "PP-X", 22, C["blue"], True)
    add_text(s, 72, 170, 330, 22, "validation-approved prior-residual", 12, C["muted"])
    add_text(s, 72, 210, 330, 80, "core + selected executor\nSun .939  RWTH .878\nMICH .751 (dual-scale)", 14, C["ink"])
    add_text(s, 72, 320, 330, 60, "paper main\nfrozen method v1", 14, C["ink"])
    add_text(s, 72, 400, 330, 40, "만능 SOTA 아님", 13, C["blue"], True)

    rect(s, 450, 110, 380, 360, C["soft_orange"], C["orange"], True)
    rect(s, 450, 110, 8, 360, C["orange"])
    nxt = rect(s, 688, 126, 118, 24, C["orange"], None, True)
    fill_shape_text(nxt, "다음", 11, C["white"], True)
    add_text(s, 474, 126, 200, 32, "PAE", 22, C["orange"], True)
    add_text(s, 474, 170, 330, 22, "식 있는 경로", 13, C["muted"])
    add_text(s, 474, 210, 330, 90, "LLM·온톨로지·source gate\n허용된 식만 실행\n이득 없으면 PP-X", 14, C["ink"])
    add_text(s, 474, 400, 330, 40, "future program", 13, C["orange"], True)

    rect(s, 852, 110, 380, 360, C["soft"], C["ink"], True)
    rect(s, 852, 110, 8, 360, C["ink"])
    later = rect(s, 1090, 126, 118, 24, C["ink"], None, True)
    fill_shape_text(later, "박사", 11, C["white"], True)
    add_text(s, 876, 126, 220, 32, "Assurance", 20, C["ink"], True)
    add_text(s, 876, 170, 330, 22, "두 경로 통합", 13, C["muted"])
    add_text(s, 876, 210, 330, 80, "믿기 / 보류 / 거절\n허가증으로 운행\nFail이면 옮기지 않음", 14, C["ink"])
    add_text(s, 876, 400, 330, 40, "이후", 13, C["ink"], True)

    add_text(
        s,
        48,
        500,
        1184,
        70,
        "PP-X는 validation-approved prior-residual framework다. executor는 근거 있을 때만 켜고, 아니면 fallback / abstention한다.",
        15,
        C["ink"],
        True,
        "center",
    )
    foot(s, p(), TOTAL)

    # PAE concept + gates
    s = blank(prs)
    head(s, "[Appendix] Future program — PAE 컨셉", "출처가 고정된 식을 컴파일하고, 적용 가능한지 검증한 뒤에만 실행한다.")
    add_text(s, 48, 84, 1184, 22, "PAE는 식 생성기가 아니다. 문헌 식 카드가 의미적으로 실행 가능하고, 검증에서 전이될 때만 켠다.", 13, C["ink"])

    inn = rect(s, 48, 116, 360, 56, C["soft"], C["ink"], True)
    fill_shape_text(inn, "문제 서술  +  데이터 스키마", 13, C["ink"], True)
    right_arrow(s, 420, 132, 28, 22, C["muted"])
    rag = rect(s, 460, 116, 360, 56, C["soft"], C["ink"], True)
    fill_shape_text(rag, "RAG   문헌 식 카드 · 인용만 검색", 13, C["ink"], True)
    add_text(s, 840, 124, 392, 40, "실행 권한 없음. 후보만 줄인다.", 12, C["muted"])

    down_arrow(s, 624, 178, 28, 14, C["muted"])

    g1 = rect(s, 48, 198, 360, 118, C["soft_blue"], C["blue"], True)
    rect(s, 48, 198, 8, 118, C["blue"])
    add_text(s, 68, 206, 320, 22, "LLM 게이트", 15, C["blue"], True)
    add_text(s, 68, 232, 320, 72, "카드 ID와 역할 결합만 제안.\n식·상수·인용·고장경계를\n만들지 못한다.", 12, C["ink"])
    right_arrow(s, 420, 240, 28, 22, C["muted"])

    g2 = rect(s, 460, 198, 360, 118, C["soft_orange"], C["orange"], True)
    rect(s, 460, 198, 8, 118, C["orange"])
    add_text(s, 480, 206, 320, 22, "온톨로지 · 식 판별", 15, C["orange"], True)
    add_text(s, 480, 232, 320, 72, "target · 메커니즘 · 역할 · 단위\n경계 · 출처 · learnable slot.\n하나라도 깨지면 실행 불가.", 12, C["ink"])
    right_arrow(s, 832, 240, 28, 22, C["muted"])

    g3 = rect(s, 872, 198, 360, 118, C["soft"], C["ink"], True)
    rect(s, 872, 198, 8, 118, C["ink"])
    add_text(s, 892, 206, 320, 22, "Source 게이트", 15, C["ink"], True)
    add_text(s, 892, 232, 320, 72, "검증 holdout에서 식 경로가\nPP-X보다 나을 때만 통과.\n안전 증명이 아니다.", 12, C["ink"])

    down_arrow(s, 624, 322, 28, 14, C["muted"])
    dec = rect(s, 310, 340, 660, 40, C["ink"], None, True)
    fill_shape_text(dec, "식이 컴파일되고, 검증에서 이기는가?", 15, C["white"], True)

    vline(s, 420, 380, 396, C["rule"])
    vline(s, 640, 380, 396, C["rule"])
    vline(s, 860, 380, 396, C["rule"])
    hline(s, 420, 860, 396, C["rule"])
    down_arrow(s, 404, 396, 28, 14, C["orange"])
    down_arrow(s, 624, 396, 28, 14, C["blue"])
    down_arrow(s, 844, 396, 28, 14, C["muted"])
    add_text(s, 300, 394, 90, 16, "YES", 11, C["orange"], True, "right")
    add_text(s, 880, 394, 90, 16, "NO", 11, C["blue"], True, "left")

    o1 = rect(s, 48, 418, 360, 100, C["soft_orange"], C["orange"], True)
    add_text(s, 64, 426, 328, 22, "PAE 실행", 14, C["orange"], True)
    add_text(s, 64, 452, 328, 56, "허용 계수만 학습.\n현재 상태 → 경계 적분. RUL=0.", 12, C["ink"])
    o2 = rect(s, 460, 418, 360, 100, C["soft_blue"], C["blue"], True)
    add_text(s, 476, 426, 328, 22, "PP-X로 되돌림", 14, C["blue"], True)
    add_text(s, 476, 452, 328, 56, "식 계약이 안 닫히거나\n검증에서 이득이 없을 때.", 12, C["ink"])
    o3 = rect(s, 872, 418, 360, 100, C["soft"], C["ink"], True)
    add_text(s, 888, 426, 328, 22, "거절", 14, C["ink"], True)
    add_text(s, 888, 452, 328, 56, "식도 없고 PP-X도 불안정하면\n예측하지 않는다.", 12, C["ink"])

    add_text(s, 48, 532, 1184, 28, "검사 항목   상태 · 고장 메커니즘 · 필수 변수 · 단위 · 경계 · 출처.   데이터셋 이름으로 경로를 고르지 않는다.", 12, C["ink"])
    add_text(s, 48, 564, 1184, 24, "자유 LLM 식 생성은 성공이 아니라 계약 위반이다. PAE는 future program이며 PP-X paper-main 숫자와 섞지 않는다.", 12, C["muted"])
    foot(s, p(), TOTAL)

    # PAE feasibility
    s = blank(prs)
    head(s, "[Appendix] 후속 — PAE 피저빌리티", "게이트가 왜 필요한지, 식이 맞을 때와 경계만 있을 때를 갈라 본다.")
    add_table(
        s,
        48,
        96,
        1184,
        240,
        ["자료", "넣은 식", "게이트 읽기", "PAE", "PP-X"],
        [
            ["알루미늄 균열", "Paris + 49.8 mm 파단", "식·경계·역할이 닫힘 → 실행", "0.969", "0.888"],
            ["미시간 배터리", "80% EOL 경계만", "메커니즘 식 없음 → 되돌려야 함", "0.635", "0.826"],
            ["NASA 실험셀", "용량 속도 quotient", "셀마다 regime이 달라 계약 약함", "0.492", "0.741"],
        ],
        font_size=13,
    )
    add_table(
        s,
        48,
        360,
        1184,
        180,
        ["이미 본 것", "아직 아닌 것"],
        [
            ["식이 맞으면 PAE가 PP-X를 이김 (Virkler +0.081)", "RAG·LLM·온톨로지를 주표에 넣지 않음"],
            ["경계만 알면 PAE를 켜면 진다 (MICH · NASA)", "source gate의 광범위 일반화"],
            ["동일-split, 오늘 포트폴리오와 분리", "LLM이 물리법칙을 발견한다는 주장"],
        ],
        font_size=13,
    )
    add_text(s, 48, 556, 1184, 36, "결론  식이 맞으면 PAE, 경계만 있으면 PP-X. 미시간 0.826은 raw-cycle 맞비교(주표 MICH 0.751).", 13, C["ink"])
    add_text(s, 48, 592, 1184, 24, "게이트는 성능을 올리는 부품이 아니라, 틀린 식을 실행하지 못하게 막는 장치다.", 12, C["muted"])
    foot(s, p(), TOTAL)

    # Close
    s = blank(prs)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    bg = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C["white"]
    bg.line.fill.background()
    frame = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, px(40), px(32), px(1200), px(656))
    frame.fill.background()
    frame.line.color.rgb = RGBColor(0x22, 0x22, 0x22)
    frame.line.width = Pt(1.25)
    add_text(s, 100, 200, 1080, 48, "감사합니다", 34, C["ink"], True, "center")
    add_text(s, 100, 270, 1080, 36, "Q & A", 22, C["blue"], True, "center")
    add_text(
        s,
        140,
        360,
        1000,
        110,
        "METHOD  Declare typed contract → Learn prior-centered residual → Approve / Decline → exact fallback / abstention\n"
        "EVIDENCE A/B  mixed strongest retrospective 9/9, p=.00390625 · equal-budget retrospective 8/9, p=.0391\n"
        "LIMITATION  DS03 prospective route-selection PASS · predictive-superiority FAIL\n"
        "SCOPE  strict-tail · unit-disjoint · val-only · PAE는 future 한 줄이며 universal SOTA를 주장하지 않음",
        14,
        C["ink"],
        False,
        "center",
    )
    add_text(s, 100, 500, 1080, 28, "질문 받겠습니다.", 16, C["muted"], False, "center")
    add_text(s, 100, 560, 1080, 22, "SPS Lab  ·  박사과정 박진서", 13, C["muted"], False, "center")
    p()

    # Appendix — backup after Q&A
    s = blank(prs)
    head(s, "[Appendix] 이게 무슨 데이터인가", "백업. 본 발표는 Q&A에서 끝낸다.")
    add_table(
        s,
        28,
        88,
        1224,
        540,
        ["표기", "물건", "한 줄"],
        [
            ["HUST", "화중과대 리튬 배터리", "충전 방법(프로토콜)이 다른 셀. 중국 랩."],
            ["Sunwoda", "선우다 상용 리튬셀", "공장 셀. 학습에 안 넣은 다른 셀로 시험."],
            ["RWTH", "아헨공대 리튬 배터리", "독일 랩 셀. 역시 처음 보는 셀로 시험."],
            ["MICH", "미시간대 리튬 배터리", "건강↔수명 관계가 학습 때와 달라진 셀."],
            ["MATR19", "MIT 2019 수명 벤치", "노트북/그리드용 셀. 고전 벤치마크."],
            ["MATRb2", "MIT 배치 2", "같은 랩, 다른 생산 로트. 시험 9셀."],
            ["NASA", "NASA PCoE 실험용 18650", "우주선/로켓이 아님. 용량 70%까지 남은 횟수."],
            ["Virkler", "알루미늄 피로균열", "판에 금이 얼마나 남았나. 배터리 아님."],
            ["N-CMAPSS", "항공기 엔진 시뮬", "NASA가 만든 디지털 터보팬. 실제 비행 기록 아님."],
        ],
        font_size=12,
    )
    add_text(s, 28, 638, 1224, 24, "NASA가 두 개다.  위는 배터리 셀, 아래 엔진은 시뮬레이터.", 12, C["muted"])
    foot(s, p(), TOTAL)

    s = blank(prs)
    head(s, "[Appendix] 누구로 나누고, 뭘 보고, 뭘 맞추나", "미지 = 학습 때 이름조차 안 본 셀/시편/엔진")
    add_table(
        s,
        20,
        86,
        1240,
        520,
        ["셋", "학습", "검증 / 시험", "보고 (X)", "맞추는 것 (Y)"],
        [
            ["화중 배터리", "충전법 1–6, 아직 건강한 구간", "충전법 7–8 / 9–10, 용량이 더 떨어진 구간", "지금까지의 용량·떨어지는 속도", "수명까지 남은 사이클"],
            ["선우다 · 아헨 · 미시간", "다른 셀의 건강한 구간", "처음 보는 셀의 말기 구간", "지금 건강, 최근 속도, 고장까지 여유", "수명까지 남은 사이클"],
            ["MIT 2019 / 배치2", "일부 셀의 건강한 구간 (배치2는 30셀)", "나머지 셀의 말기 (배치2는 9+9셀)", "지금까지의 용량 이력", "수명까지 남은 사이클"],
            ["NASA 실험셀", "일부 셀, 용량이 아직 높은 구간", "다른 셀, 용량이 학습 최저보다 낮은 구간", "용량, 속도, 몇 번째 사이클인가", "용량 70% 될 때까지 남은 횟수"],
            ["알루미늄 균열", "다른 시편의 짧은 금", "처음 보는 시편의 긴 금", "금 길이, 자라는 속도, 하중", "쪼개질 때까지 남은 반복"],
            ["항공기 엔진", "본 적 있는 비행조건의 엔진", "처음 보는 비행조건의 엔진", "센서(조건으로 나눈 값), 운전모드", "엔진이 버틸 남은 시간"],
        ],
        font_size=11,
    )
    add_text(s, 20, 618, 1240, 28, "같은 셀을 시간만 잘라 뒤를 맞추지 않는다.  시험 셀의 미래·최종 수명은 X에 넣지 않는다.", 12, C["muted"])
    foot(s, p(), TOTAL)

    # Put every explicitly marked backup after Q&A, then enforce the main-story
    # order independently of source build order.
    sld_ids = list(prs.slides._sldIdLst)
    appendix_ids = []
    main_ids = []
    for sld_id, slide in zip(sld_ids, prs.slides):
        text = "\n".join(sh.text for sh in slide.shapes if hasattr(sh, "text_frame"))
        (appendix_ids if "[Appendix]" in text else main_ids).append(sld_id)
    desired_main_titles = [
        "PP-X: Prior-Transferability Adaptive Extrapolation",
        "문제 · 연구 배경",
        "연구 루트",
        "PP-X 개요",
        "왜 PP-X가 필요한가",
        "본 연구의 기여",
        "기존 방법과의 차이",
        "결과의 해석 범위",
        "방법 — Algorithm 1: Declare → Learn → Approve → Decline",
    ]
    title_to_id = {}
    for sld_id, slide in zip(sld_ids, prs.slides):
        slide_texts = {
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        }
        for title in desired_main_titles:
            if title in slide_texts:
                title_to_id[title] = sld_id
    ordered_main = [title_to_id[title] for title in desired_main_titles]
    main_ids = ordered_main + [sld_id for sld_id in main_ids if sld_id not in ordered_main]

    for sld_id in sld_ids:
        prs.slides._sldIdLst.remove(sld_id)
    for sld_id in main_ids + appendix_ids:
        prs.slides._sldIdLst.append(sld_id)

    # Reordering must not leave historical build-order page numbers in the
    # appendix. Preserve the original text formatting and replace only the
    # numeric footer token.
    for index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            token = shape.text.strip()
            parts = token.split("/")
            if len(parts) != 2 or not all(part.isdigit() for part in parts):
                continue
            paragraph = shape.text_frame.paragraphs[0]
            if paragraph.runs:
                paragraph.runs[0].text = f"{index}/{TOTAL}"

    assert n == TOTAL, f"TOTAL mismatch: built {n}, expected {TOTAL}"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Saved {OUT} ({n} slides)")


if __name__ == "__main__":
    build()
